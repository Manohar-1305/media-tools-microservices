
import io
import logging
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
)
from xml.sax.saxutils import escape

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".docx"}


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html")


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status": "healthy",
        "service": "word-to-pdf"
    }), 200


@app.route("/convert", methods=["POST"])
def convert():
    uploaded_file = request.files.get("file")

    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"error": "Please select a Word document."}), 400

    filename = Path(uploaded_file.filename).name

    if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
        return jsonify({
            "error": "This Python-only converter supports DOCX files."
        }), 400

    try:
        document = Document(io.BytesIO(uploaded_file.read()))

        pdf_buffer = io.BytesIO()
        styles = getSampleStyleSheet()

        pdf = SimpleDocTemplate(
            pdf_buffer,
            pagesize=A4,
            rightMargin=50,
            leftMargin=50,
            topMargin=50,
            bottomMargin=50,
            title=Path(filename).stem,
        )

        story = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if not text:
                story.append(Spacer(1, 8))
                continue

            safe_text = escape(text).replace("\n", "<br/>")
            style = styles["Title"] if paragraph.style.name == "Title" else styles["Normal"]

            story.append(Paragraph(safe_text, style))
            story.append(Spacer(1, 6))

        for table in document.tables:
            for row in table.rows:
                cells = [
                    escape(cell.text.strip()).replace("\n", "<br/>")
                    for cell in row.cells
                ]
                story.append(Paragraph(" | ".join(cells), styles["Normal"]))
                story.append(Spacer(1, 4))

        if not story:
            story.append(Paragraph("The document contains no text.", styles["Normal"]))

        pdf.build(story)
        pdf_buffer.seek(0)

        return send_file(
            pdf_buffer,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"{Path(filename).stem}.pdf",
        )

    except Exception:
        logger.exception("DOCX-to-PDF conversion failed")
        return jsonify({
            "error": "Unable to convert this DOCX document."
        }), 422


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "Maximum file size is 25 MB."}), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5011)
