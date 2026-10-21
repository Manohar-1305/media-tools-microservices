
import io
import logging
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file
from pypdf import PdfReader
from docx import Document

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

@app.route("/", methods=["GET"])
def index():
    return render_template("index.html")


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "status": "healthy",
        "service": "pdf-to-word"
    }), 200


@app.route("/convert", methods=["POST"])
def convert():
    uploaded_file = request.files.get("file")

    if not uploaded_file or not uploaded_file.filename:
        return jsonify({"error": "Please select a PDF file."}), 400

    filename = Path(uploaded_file.filename).name

    if Path(filename).suffix.lower() != ".pdf":
        return jsonify({"error": "Only PDF files are supported."}), 400

    try:
        pdf_data = uploaded_file.read()

        if not pdf_data:
            return jsonify({"error": "The selected PDF is empty."}), 400

        reader = PdfReader(io.BytesIO(pdf_data))

        if reader.is_encrypted:
            return jsonify({
                "error": "Password-protected PDFs are not supported."
            }), 422

        document = Document()
        document.add_heading(Path(filename).stem, level=1)

        extracted_pages = 0

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text()

            if text and text.strip():
                extracted_pages += 1

                for paragraph_text in text.splitlines():
                    if paragraph_text.strip():
                        document.add_paragraph(paragraph_text)

            if page_number < len(reader.pages):
                document.add_page_break()

        if extracted_pages == 0:
            return jsonify({
                "error": (
                    "No extractable text was found. "
                    "This PDF may contain scanned images and require OCR."
                )
            }), 422

        output = io.BytesIO()
        document.save(output)
        output.seek(0)

        logger.info(
            "Converted PDF to DOCX: filename=%s pages=%s",
            filename,
            len(reader.pages),
        )

        return send_file(
            output,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            as_attachment=True,
            download_name=f"{Path(filename).stem}.docx",
        )

    except Exception:
        logger.exception("PDF-to-Word conversion failed")
        return jsonify({
            "error": "Unable to convert this PDF. Check that it is valid and readable."
        }), 422


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "Maximum file size is 25 MB."}), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5012)
